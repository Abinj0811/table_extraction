import os
import openpyxl
import pandas as pd
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image
from PIL import Image as PILImage

def check_directory_exists(directory):
    """Check if the specified directory exists."""
    return os.path.exists(directory)

def get_sorted_excel_dirs(pdf_img_dir):
    """Get sorted directories containing Excel files."""
    excel_dirs = [os.path.join(pdf_img_dir, folder) for folder in os.listdir(pdf_img_dir) if os.path.isdir(os.path.join(pdf_img_dir, folder))]
    sorted_excel_dirs = sorted(
        excel_dirs,
        key=lambda x: int(os.path.basename(x).split('_')[-1]) if '_' in os.path.basename(x) else float('inf')
    )
    return sorted_excel_dirs

def create_workbook():
    """Create a new Excel workbook and remove the default sheet."""
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    return workbook

def process_excel_files(sheet, folder):
    """Process each Excel file and copy its content to the specified sheet."""
    excel_files = [f for f in os.listdir(folder) if f.lower().endswith('.xlsx')]
    for excel_file in excel_files:
        excel_file_path = os.path.join(folder, excel_file)
        try:
            input_workbook = openpyxl.load_workbook(excel_file_path, data_only=True)
            for input_sheet_name in input_workbook.sheetnames:
                input_sheet = input_workbook[input_sheet_name]
                data = input_sheet.values
                df = pd.DataFrame(data)

                start_row = sheet.max_row + 1
                sheet.append([f"Sheet: {input_sheet_name}"])
                sheet.append([])

                for row in dataframe_to_rows(df, index=False, header=False):
                    sheet.append(row)

                # Add three empty rows for separation between tables
                for _ in range(3):
                    sheet.append([])

                # Add corresponding images
                add_images_to_sheet(sheet, folder, input_sheet_name, start_row)

        except Exception as e:
            print(f"Error processing Excel file '{excel_file}': {e}")


def add_images_to_sheet(sheet, folder, input_sheet_name, start_row):
    """Add images to the specified sheet based on the input sheet name."""
    image_files = [f for f in os.listdir(folder) if f.lower().endswith('.png')]
    for image_file in image_files:
        if input_sheet_name in image_file:  # Check if the image corresponds to the sheet
            img_path = os.path.join(folder, image_file)
            resize_and_save_image(img_path)

            img = Image(img_path)
            last_column = sheet.max_column
            img_anchor = f"{get_column_letter(last_column + 3)}{start_row}"
            img.anchor = img_anchor
            sheet.add_image(img)
            break  # Assuming one image per sheet

def resize_and_save_image(img_path):
    """Resize the image to 50% of its original size and save it."""
    with PILImage.open(img_path) as img:
        # Get the original size
        original_size = img.size
        # Calculate new size as 50% of the original size
        new_size = (int(original_size[0] * 0.5), int(original_size[1] * 0.5))
        # Resize the image
        img = img.resize(new_size, PILImage.LANCZOS)
        # Save the resized image
        img.save(img_path)


def save_workbook(workbook, output_excel_path):
    """Save the workbook to the specified output path."""
    workbook.save(output_excel_path)
    print(f"Successfully created Excel file: {output_excel_path}")

def main():
    pdf_img_dir = 'images/K Ship_Bocom(S-1940)_ME Mach_Final dwg_24.04.30'
    output_excel_name = os.path.basename(pdf_img_dir) + '.xlsx'
    output_excel_path = os.path.join('out_excel', output_excel_name)

    if check_directory_exists(pdf_img_dir):
        sorted_excel_dirs = get_sorted_excel_dirs(pdf_img_dir)
        workbook = create_workbook()

        for folder in sorted_excel_dirs:
            folder_name = os.path.basename(folder)
            sheet = workbook.create_sheet(title=folder_name)
            process_excel_files(sheet, folder)

        save_workbook(workbook, output_excel_path)
    else:
        print("The specified directory does not exist.")

if __name__ == "__main__":
    main()